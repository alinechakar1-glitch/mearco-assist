"""MEARCO ONE pilot API. Payment gateway/voice/WhatsApp are NOT active integrations."""
import os, io, secrets, datetime as dt
from decimal import Decimal
from typing import Literal
from fastapi import FastAPI, HTTPException, Depends, UploadFile, File, Header
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import HTMLResponse
from pydantic import BaseModel, Field
from sqlalchemy import create_engine, Column, Integer, String, ForeignKey, Numeric, Date, DateTime, Text, select, func
from sqlalchemy.orm import declarative_base, sessionmaker, Session
from passlib.context import CryptContext
import jwt
from openpyxl import load_workbook

DATABASE_URL=os.getenv('DATABASE_URL','sqlite:///./mearco.db')
if DATABASE_URL.startswith('postgres://'): DATABASE_URL=DATABASE_URL.replace('postgres://','postgresql+psycopg://',1)
if DATABASE_URL.startswith('postgresql://'): DATABASE_URL=DATABASE_URL.replace('postgresql://','postgresql+psycopg://',1)
SECRET=os.getenv('JWT_SECRET') or os.getenv('SESSION_SECRET','')
if not SECRET: SECRET=secrets.token_urlsafe(48); print('WARNING: ephemeral JWT secret; set JWT_SECRET in production')
engine=create_engine(DATABASE_URL,connect_args={'check_same_thread':False} if DATABASE_URL.startswith('sqlite') else {},pool_pre_ping=True)
SessionLocal=sessionmaker(bind=engine)
Base=declarative_base()
crypt=CryptContext(schemes=['bcrypt'],deprecated='auto')
class User(Base):
 __tablename__='users'
 id=Column(Integer,primary_key=True); username=Column(String(120),unique=True,index=True); pw=Column(String(255)); role=Column(String(30)); party_code=Column(String(90),nullable=True)
class Policy(Base):
 __tablename__='policies'
 id=Column(Integer,primary_key=True); number=Column(String(100),unique=True,index=True); customer_code=Column(String(90),index=True); customer_name=Column(String(200)); broker_code=Column(String(90),index=True); product=Column(String(60)); plate=Column(String(100)); start=Column(Date); end=Column(Date); towing_total=Column(Integer,default=0); premium=Column(Numeric(14,2),default=0); due=Column(Numeric(14,2),default=0); status=Column(String(30),default='active'); pdf_url=Column(String(1000),default='')
class Case(Base):
 __tablename__='cases'
 id=Column(Integer,primary_key=True); policy_id=Column(Integer,ForeignKey('policies.id')); category=Column(String(40)); description=Column(Text); status=Column(String(40),default='new'); provider_id=Column(Integer,ForeignKey('providers.id'),nullable=True); cost=Column(Numeric(14,2),default=0); created=Column(DateTime,default=dt.datetime.utcnow); address=Column(String(300),default='')
class Provider(Base):
 __tablename__='providers'
 id=Column(Integer,primary_key=True); name=Column(String(180)); kind=Column(String(40)); area=Column(String(100)); phone=Column(String(80)); active=Column(String(20),default='yes')
class Payment(Base):
 __tablename__='payments'
 id=Column(Integer,primary_key=True); policy_id=Column(Integer,ForeignKey('policies.id')); amount=Column(Numeric(14,2)); status=Column(String(35),default='pending'); reference=Column(String(120),unique=True); checkout_url=Column(String(1000),default=''); created=Column(DateTime,default=dt.datetime.utcnow)
class Audit(Base):
 __tablename__='audit'
 id=Column(Integer,primary_key=True); actor=Column(String(120)); action=Column(String(120)); detail=Column(Text); at=Column(DateTime,default=dt.datetime.utcnow)
Base.metadata.create_all(engine)
app=FastAPI(title='MEARCO Unified Pilot',version='1.0.0')

@app.on_event('startup')
def ensure_admin_account():
 """Provision admin from Render environment, no shell required.

 Set ADMIN_PASSWORD in Render and optionally ADMIN_USERNAME (default: admin).
 This also updates the existing admin password if the configured value changes.
 Other users and their passwords are never touched.
 """
 admin_username = os.getenv('ADMIN_USERNAME', 'admin').strip()
 admin_password = os.getenv('ADMIN_PASSWORD', '')
 if not admin_password:
  print('ADMIN SETUP: ADMIN_PASSWORD is missing; administrator auto-setup skipped.')
  return
 if not admin_username or len(admin_username) > 120:
  raise RuntimeError('ADMIN_USERNAME must be between 1 and 120 characters.')
 if len(admin_password) < 12 or len(admin_password.encode('utf-8')) > 72:
  raise RuntimeError('ADMIN_PASSWORD must be at least 12 characters and at most 72 UTF-8 bytes (bcrypt limit).')
 with SessionLocal() as session:
  existing = session.scalar(select(User).where(User.username == admin_username))
  if existing is None:
   session.add(User(username=admin_username, pw=crypt.hash(admin_password), role='admin'))
   session.commit()
   print('ADMIN SETUP: Administrator account created successfully.')
  else:
   password_matches = crypt.verify(admin_password, existing.pw)
   if not password_matches or existing.role != 'admin':
    if not password_matches:
     existing.pw = crypt.hash(admin_password)
    existing.role = 'admin'
    session.commit()
   print('ADMIN SETUP: Administrator account ready.')

origins=[x.strip() for x in os.getenv('CORS_ORIGINS','http://localhost:8081,http://localhost:19006').split(',') if x.strip()]
app.add_middleware(CORSMiddleware,allow_origins=origins,allow_methods=['GET','POST','PATCH'],allow_headers=['*'])
def db():
 with SessionLocal() as s: yield s
def audit(s,u,act,detail): s.add(Audit(actor=u.username,action=act,detail=str(detail)))
def user(authorization: str=Header(default=''),s:Session=Depends(db)):
 try:
  token=authorization.removeprefix('Bearer ').strip(); data=jwt.decode(token,SECRET,algorithms=['HS256']); u=s.get(User,int(data['sub']))
  if not u: raise ValueError('missing account')
  return u
 except Exception: raise HTTPException(401,'Sign in required')
def roles(u,*allowed):
 if u.role not in allowed: raise HTTPException(403,'Access denied')
def visible(p,u):
 return u.role in ('admin','staff') or (u.role=='broker' and p.broker_code==u.party_code) or (u.role=='customer' and p.customer_code==u.party_code)
def policy_access(s,u,pid):
 p=s.get(Policy,pid)
 if not p: raise HTTPException(404,'Policy not found')
 if not visible(p,u): raise HTTPException(403,'Not your policy')
 return p
def pol(p): return {k:(float(getattr(p,k)) if k in ('premium','due') else getattr(p,k).isoformat() if k in ('start','end') and getattr(p,k) else getattr(p,k)) for k in ('id','number','customer_code','customer_name','broker_code','product','plate','start','end','towing_total','premium','due','status','pdf_url')}
def case_obj(c): return {'id':c.id,'policy_id':c.policy_id,'category':c.category,'description':c.description,'status':c.status,'provider_id':c.provider_id,'cost':float(c.cost or 0),'address':c.address,'created':c.created.isoformat() if c.created else None}
class Login(BaseModel): username:str; password:str
class CaseIn(BaseModel): policy_id:int; category:Literal['towing','motor_claim','assistance','garage']; description:str=Field(min_length=5,max_length=3000); address:str=''
class CaseUpdate(BaseModel): status:Literal['new','triaged','dispatched','in_progress','awaiting_approval','approved','declined','closed']; provider_id:int|None=None; cost:float|None=None
class ProviderIn(BaseModel): name:str; kind:Literal['towing','garage']; area:str=''; phone:str=''
class PaymentIn(BaseModel): policy_id:int; amount:float=Field(gt=0)
class PaymentUpdate(BaseModel): status:Literal['paid','failed','cancelled']
class BotQuery(BaseModel): policy_id:int|None=None; question:str=Field(min_length=1,max_length=1500)
@app.get('/health')
def health(): return {'ok':True,'mode':'pilot','real_payments':False,'live_voice':False,'live_whatsapp':False}
@app.post('/auth/login')
def login(body:Login,s:Session=Depends(db)):
 u=s.scalar(select(User).where(User.username==body.username))
 if not u or not crypt.verify(body.password,u.pw): raise HTTPException(401,'Invalid credentials')
 token=jwt.encode({'sub':str(u.id),'exp':dt.datetime.now(dt.timezone.utc)+dt.timedelta(hours=8)},SECRET,algorithm='HS256')
 return {'access_token':token,'role':u.role,'username':u.username,'party_code':u.party_code}
@app.get('/policies')
def policies(u:User=Depends(user),s:Session=Depends(db)):
 q=select(Policy)
 if u.role=='broker': q=q.where(Policy.broker_code==u.party_code)
 elif u.role=='customer': q=q.where(Policy.customer_code==u.party_code)
 elif u.role not in ('admin','staff'): raise HTTPException(403)
 return [pol(p) for p in s.scalars(q.order_by(Policy.id.desc())).all()]
@app.get('/policies/{pid}')
def policy(pid:int,u:User=Depends(user),s:Session=Depends(db)): return pol(policy_access(s,u,pid))
@app.get('/policies/{pid}/towing-balance')
def towing_balance(pid:int,u:User=Depends(user),s:Session=Depends(db)):
 p=policy_access(s,u,pid)
 used=s.scalar(select(func.count(Case.id)).where(Case.policy_id==pid,Case.category=='towing',Case.status!='declined')) or 0
 valid=bool(p.status=='active' and p.start<=dt.date.today()<=p.end)
 return {'entitled':p.towing_total,'used':used,'remaining':max(0,p.towing_total-used),'policy_in_force':valid,'note':'Manual eligibility review required for exceptions and policy terms'}
@app.post('/cases')
def open_case(v:CaseIn,u:User=Depends(user),s:Session=Depends(db)):
 p=policy_access(s,u,v.policy_id)
 if u.role=='provider':raise HTTPException(403)
 if v.category=='towing':
  used=s.scalar(select(func.count(Case.id)).where(Case.policy_id==p.id,Case.category=='towing',Case.status!='declined')) or 0
  if p.product.lower()!='motor' or p.status!='active' or not(p.start<=dt.date.today()<=p.end) or used>=p.towing_total: raise HTTPException(409,'No confirmed towing entitlement; call MEARCO for manual review')
 c=Case(policy_id=p.id,category=v.category,description=v.description,address=v.address);s.add(c);s.flush();audit(s,u,'create_case',c.id);s.commit();return case_obj(c)
@app.get('/cases')
def cases(u:User=Depends(user),s:Session=Depends(db)):
 q=select(Case).join(Policy)
 if u.role=='customer':q=q.where(Policy.customer_code==u.party_code)
 elif u.role=='broker':q=q.where(Policy.broker_code==u.party_code)
 elif u.role=='provider':q=q.where(Case.provider_id==s.scalar(select(Provider.id).where(Provider.name==u.party_code)))
 elif u.role not in ('admin','staff'):raise HTTPException(403)
 return [case_obj(c) for c in s.scalars(q.order_by(Case.id.desc())).all()]
@app.patch('/cases/{cid}')
def update_case(cid:int,v:CaseUpdate,u:User=Depends(user),s:Session=Depends(db)):
 roles(u,'admin','staff');c=s.get(Case,cid)
 if not c:raise HTTPException(404)
 if v.provider_id is not None and not s.get(Provider,v.provider_id):raise HTTPException(404,'Provider missing')
 c.status=v.status
 if v.provider_id is not None:c.provider_id=v.provider_id
 if v.cost is not None:c.cost=Decimal(str(v.cost))
 audit(s,u,'update_case',f'{cid}:{v.status}');s.commit();return case_obj(c)
@app.get('/providers')
def providers(u:User=Depends(user),s:Session=Depends(db)):
 return [{'id':p.id,'name':p.name,'kind':p.kind,'area':p.area,'phone':p.phone} for p in s.scalars(select(Provider).where(Provider.active=='yes')).all()]
@app.post('/providers')
def create_provider(v:ProviderIn,u:User=Depends(user),s:Session=Depends(db)):
 roles(u,'admin','staff');p=Provider(**v.model_dump());s.add(p);s.flush();audit(s,u,'create_provider',p.id);s.commit();return {'id':p.id}
@app.get('/balances')
def balances(u:User=Depends(user),s:Session=Depends(db)):
 q=select(Policy).where(Policy.due>0)
 if u.role=='broker':q=q.where(Policy.broker_code==u.party_code)
 elif u.role=='customer':q=q.where(Policy.customer_code==u.party_code)
 elif u.role not in ('admin','staff'):raise HTTPException(403)
 return [{'policy_id':p.id,'number':p.number,'customer_name':p.customer_name,'due':float(p.due)} for p in s.scalars(q).all()]
@app.post('/payments/requests')
def payment_request(v:PaymentIn,u:User=Depends(user),s:Session=Depends(db)):
 p=policy_access(s,u,v.policy_id)
 if Decimal(str(v.amount))>p.due:raise HTTPException(409,'Payment exceeds policy balance')
 ref='MEARCO-'+secrets.token_hex(6).upper()
 pay=Payment(policy_id=p.id,amount=Decimal(str(v.amount)),reference=ref);s.add(pay);audit(s,u,'payment_request',ref);s.commit()
 return {'reference':ref,'status':'pending','checkout_url':None,'instructions':'No payment link until Whish/Monty or bank gateway is configured and verified; never mark paid from customer input.'}
@app.get('/payments')
def list_payments(u:User=Depends(user),s:Session=Depends(db)):
 q=select(Payment).join(Policy)
 if u.role=='broker':q=q.where(Policy.broker_code==u.party_code)
 elif u.role=='customer':q=q.where(Policy.customer_code==u.party_code)
 elif u.role not in ('admin','staff'):raise HTTPException(403)
 return [{'reference':x.reference,'policy_id':x.policy_id,'amount':float(x.amount),'status':x.status,'checkout_url':x.checkout_url or None} for x in s.scalars(q).all()]
@app.post('/payments/{ref}/record-manual')
def manual_record(ref:str,v:PaymentUpdate,u:User=Depends(user),s:Session=Depends(db)):
 roles(u,'admin');x=s.scalar(select(Payment).where(Payment.reference==ref))
 if not x:raise HTTPException(404)
 if x.status!='pending':raise HTTPException(409,'Already processed')
 if v.status=='paid':
  p=s.get(Policy,x.policy_id)
  if x.amount>p.due:raise HTTPException(409,'Outstanding amount has changed')
  p.due=p.due-x.amount
 x.status=v.status;audit(s,u,'manual_payment_record',f'{ref}:{v.status}');s.commit();return {'reference':ref,'status':x.status,'warning':'Pilot manual entry only; verify bank settlement independently'}
@app.post('/assistant/message')
def assistant(v:BotQuery,u:User=Depends(user),s:Session=Depends(db)):
 q=v.question.lower()
 if v.policy_id: policy_access(s,u,v.policy_id)
 if any(w in q for w in ('accident','emergency','injured','حادث','مصاب','tow','towing','سطحة')):
  ans='For an emergency, call local emergency services if anyone is injured. Contact MEARCO hotline immediately. I can help you create an assistance request, but dispatch requires confirmation.'
 elif any(w in q for w in ('pay','balance','unpaid','دفع','متوجب')):ans='View balances in the app. Payment links are not active yet; contact MEARCO for an approved payment method.'
 elif any(w in q for w in ('policy','renew','وثيقة','بوليصة')):ans='Your policy portfolio and expiry dates are available under Policies. Renewal requires MEARCO approval.'
 else:ans='I can assist with policies, balances, towing and claims. For a coverage decision or urgent situation I will refer you to a MEARCO representative.'
 return {'answer':ans,'handoff_required':any(w in q for w in ('emergency','injured','حادث','مصاب','complaint','رفض','declined')),'automated':'rules_only','note':'Not a live phone/LLM assistant'}
@app.post('/admin/import/policies')
async def import_policies(file:UploadFile=File(...),u:User=Depends(user),s:Session=Depends(db)):
 roles(u,'admin');raw=await file.read()
 if len(raw)>5_000_000:raise HTTPException(413,'Maximum 5 MB')
 try: w=load_workbook(io.BytesIO(raw),read_only=True,data_only=True); sh=w['Policies']; rows=sh.iter_rows(values_only=True); heads=[str(x).strip() for x in next(rows)]
 except Exception:raise HTTPException(400,'Expected an .xlsx workbook with a Policies sheet')
 required=['number','customer_code','customer_name','broker_code','product','plate','start','end','towing_total','premium','due','status']
 if heads!=required:raise HTTPException(400,'Columns must exactly match: '+','.join(required))
 count=0;errors=[]
 for line,vals in enumerate(rows,2):
  if not any(x is not None for x in vals):continue
  try:
   d=dict(zip(heads,vals)); number=str(d['number']).strip()
   if not number:raise ValueError('blank number')
   p=s.scalar(select(Policy).where(Policy.number==number)) or Policy(number=number)
   for k in ('customer_code','customer_name','broker_code','product','plate','status'):setattr(p,k,str(d[k] or ''))
   for k in ('start','end'):
    value=d[k];setattr(p,k,value.date() if isinstance(value,dt.datetime) else value if isinstance(value,dt.date) else dt.date.fromisoformat(str(value)))
   p.towing_total=int(d['towing_total'] or 0);p.premium=Decimal(str(d['premium'] or 0));p.due=Decimal(str(d['due'] or 0))
   if p.towing_total<0 or p.due<0 or p.premium<0 or p.due>p.premium:raise ValueError('invalid amounts or towing total')
   s.add(p);count+=1
  except Exception as e: errors.append({'line':line,'error':str(e)})
 if errors:s.rollback();return {'imported':0,'errors':errors[:30],'note':'No data committed; correct workbook and retry'}
 audit(s,u,'policy_import',count);s.commit();return {'imported':count,'errors':[]}
@app.get('/admin/summary')
def summary(u:User=Depends(user),s:Session=Depends(db)):
 roles(u,'admin','staff')
 return {'policies':s.scalar(select(func.count(Policy.id))),'all_cases':s.scalar(select(func.count(Case.id))),'new_cases':s.scalar(select(func.count(Case.id)).where(Case.status=='new')),'motor_vehicles':s.scalar(select(func.count(Policy.id)).where(Policy.product=='Motor')),'completed_tows':s.scalar(select(func.count(Case.id)).where(Case.category=='towing',Case.status=='closed')),'open_cases':s.scalar(select(func.count(Case.id)).where(Case.status!='closed')),'outstanding_usd':float(s.scalar(select(func.coalesce(func.sum(Policy.due),0)))),'pending_payment_requests':s.scalar(select(func.count(Payment.id)).where(Payment.status=='pending'))}
@app.get('/old-pilot-landing',response_class=HTMLResponse)
def landing():return '''<html><head><meta name="viewport" content="width=device-width,initial-scale=1"><style>body{font:16px system-ui;background:#f0f7f7;color:#123;padding:3rem;max-width:700px;margin:auto}a{color:#087f78}.card{background:white;border-radius:16px;padding:24px}</style></head><body><div class="card"><h1>MEARCO ONE • Pilot API</h1><p>Backend service is running. This is not the existing MEARCO Assist production site.</p><p><a href="/docs">Open API documentation</a> · <a href="/health">System status</a></p><p>Install the mobile project separately. Do not use real customers or issue real payment links in this pilot.</p></div></body></html>'''

from fastapi.responses import FileResponse
from pathlib import Path as _Path
@app.get('/',include_in_schema=False)
def web_app():
 return FileResponse(_Path(__file__).parent/'static'/'index.html',media_type='text/html')
@app.get('/app',include_in_schema=False)
def web_app_alias(): return web_app()
@app.get('/integrations/status')
def integration_status(u:User=Depends(user)):
 roles(u,'admin','staff')
 return {'original_assist_recreated':True,'original_v07_exact_code_merged':False,'mobile_app':'source included','payment_gateway':'not connected','voice_AI':'not connected','whatsapp':'not connected','store_release':'not published','database':'configured'}
