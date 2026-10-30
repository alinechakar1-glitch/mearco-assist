"""MEARCO Assist prototype. Do not put real personal data into this demo."""
import os, sqlite3, io, csv, secrets, hashlib, time, hmac, re
from datetime import datetime, timezone
from pathlib import Path
from contextlib import contextmanager
from fastapi import FastAPI, Request, Form, UploadFile, File, HTTPException
from fastapi.responses import HTMLResponse, RedirectResponse, JSONResponse
from fastapi.templating import Jinja2Templates
from fastapi.staticfiles import StaticFiles
from starlette.middleware.sessions import SessionMiddleware
from openpyxl import load_workbook

ROOT=Path(__file__).parent
DB_PATH=os.environ.get('DB_PATH',str(ROOT/'assist_demo.sqlite3'))
ADMIN_PASSWORD=os.environ.get('ADMIN_PASSWORD')
SESSION_SECRET=os.environ.get('SESSION_SECRET')
PROVIDER_PASSWORD=os.environ.get('PROVIDER_PASSWORD','')
PUBLIC_INTAKE_ENABLED=os.environ.get('PUBLIC_INTAKE_ENABLED','0')=='1'
PUBLIC_TRACKING_ENABLED=os.environ.get('PUBLIC_TRACKING_ENABLED','0')=='1'
TRUSTED_ORIGINS=[v.strip().rstrip('/') for v in os.environ.get('TRUSTED_ORIGINS','').split(',') if v.strip()]
if not ADMIN_PASSWORD or not SESSION_SECRET:
    raise RuntimeError('Set ADMIN_PASSWORD and SESSION_SECRET environment variables before starting')
app=FastAPI(title='MEARCO Assist 24/7 — Standalone')

@app.get('/version')
def version(): return {'version': '0.7.0', 'csrf_protection': 'session-form-token', 'mode':'TEST — not live dispatch'}

@app.middleware('http')
async def security_headers(request,call_next):
    response=await call_next(request)
    response.headers['X-Content-Type-Options']='nosniff'
    response.headers['X-Frame-Options']='DENY'
    response.headers['Referrer-Policy']='no-referrer'
    response.headers['Cache-Control']='no-store'
    return response

app.add_middleware(SessionMiddleware,secret_key=SESSION_SECRET,https_only=os.environ.get('COOKIE_SECURE','1')=='1',same_site='lax')
app.mount('/static',StaticFiles(directory=str(ROOT/'static')),name='static')
templates=Jinja2Templates(directory=str(ROOT/'templates'))

@contextmanager
def conn():
    db=sqlite3.connect(DB_PATH, timeout=20)
    db.row_factory=sqlite3.Row
    try:
        yield db
        db.commit()
    finally: db.close()

def now(): return datetime.now(timezone.utc).isoformat(timespec='seconds')

def init_db():
    with conn() as db:
        db.execute('CREATE TABLE IF NOT EXISTS policies(id INTEGER PRIMARY KEY,policy_number TEXT NOT NULL UNIQUE,insured_name TEXT,phone TEXT,plate TEXT,start_date TEXT,end_date TEXT,towing_entitlement TEXT DEFAULT "Review",status TEXT DEFAULT "Unknown")')
        db.execute('CREATE TABLE IF NOT EXISTS cases(id INTEGER PRIMARY KEY,created_at TEXT NOT NULL,client_name TEXT NOT NULL,phone TEXT,plate TEXT,location TEXT,case_type TEXT,status TEXT,notes TEXT,provider TEXT,eta TEXT)')
        db.execute('CREATE TABLE IF NOT EXISTS case_events(id INTEGER PRIMARY KEY,case_id INTEGER NOT NULL,created_at TEXT NOT NULL,actor TEXT NOT NULL,event TEXT NOT NULL)')
        db.execute('CREATE TABLE IF NOT EXISTS public_tokens(case_id INTEGER PRIMARY KEY, token_hash TEXT NOT NULL)')
        db.execute('CREATE TABLE IF NOT EXISTS provider_jobs(case_id INTEGER PRIMARY KEY, provider_note TEXT, assigned_driver TEXT, updated_at TEXT)')
init_db()

def has_table(db,name):
    return db.execute("SELECT 1 FROM sqlite_master WHERE type='table' AND name=?",(name,)).fetchone() is not None

def hash_password(password):
    salt=secrets.token_bytes(16)
    derived=hashlib.pbkdf2_hmac('sha256',password.encode(),salt,310000)
    return 'pbkdf2_sha256$310000$'+salt.hex()+'$'+derived.hex()
def verify_password(password,encoded):
    try:
        algo,count,salt_hex,expected=encoded.split('$')
        if algo!='pbkdf2_sha256':return False
        actual=hashlib.pbkdf2_hmac('sha256',password.encode(),bytes.fromhex(salt_hex),int(count))
        return hmac.compare_digest(actual,bytes.fromhex(expected))
    except (ValueError,TypeError):return False

def ensure_staff():
    with conn() as db:
        db.execute('CREATE TABLE IF NOT EXISTS staff_users(id INTEGER PRIMARY KEY,username TEXT NOT NULL UNIQUE COLLATE NOCASE,display_name TEXT NOT NULL,role TEXT NOT NULL,active INTEGER NOT NULL DEFAULT 1,password_hash TEXT NOT NULL,created_at TEXT NOT NULL,last_login TEXT)')
        db.execute('CREATE TABLE IF NOT EXISTS staff_audit(id INTEGER PRIMARY KEY,created_at TEXT NOT NULL,actor TEXT NOT NULL,action TEXT NOT NULL,subject TEXT NOT NULL)')
ensure_staff()

def authorized(req): return bool(req.session.get('signed_in')) and req.session.get('role') in ('admin','supervisor','operator')
def is_admin(req): return authorized(req) and req.session.get('role')=='admin'
def current_actor(req): return req.session.get('username','MEARCO operator')
def provider_authorized(req): return bool(req.session.get('signed_in')) and req.session.get('role')=='provider'
def new_csrf(req):
    if not req.session.get('csrf'): req.session['csrf']=secrets.token_urlsafe(24)
    return req.session['csrf']
def csrf_ok(req,token): return bool(token) and secrets.compare_digest(str(token),str(req.session.get('csrf','')))
def enforce_csrf(req,token):
    # A synchronizer CSRF token is generated on GET and required on every POST.
    # Do not compare Origin to Render's proxy-facing URL: proxies and browsers
    # may vary these headers. SameSite=Lax and a per-session secret token
    # protect the form without an unreliable deployment-specific origin check.
    if not csrf_ok(req, token):
        raise HTTPException(403, 'Invalid form token; refresh page')

def redirect(location): return RedirectResponse(location,status_code=303)
def render(req,name,**kw): return templates.TemplateResponse(req,name,{'logged_in':authorized(req),'is_admin':is_admin(req),'username':req.session.get('username',''),'provider_logged_in':provider_authorized(req),'csrf_token':new_csrf(req),'public_enabled':PUBLIC_INTAKE_ENABLED,**kw})

@app.get('/login',response_class=HTMLResponse)
def login_form(request:Request): return render(request,'login.html',error=None)
@app.post('/login')
def login(request:Request,password:str=Form(...),csrf_token:str=Form('')):
    enforce_csrf(request,csrf_token)
    if secrets.compare_digest(password,ADMIN_PASSWORD):
        request.session.clear()
        request.session.update({'signed_in':True,'role':'admin','username':'Administrator'})
        return redirect('/')
    return templates.TemplateResponse(request,'login.html',{'logged_in':False,'error':'Incorrect password'},status_code=401)
@app.post('/staff-login')
def staff_login(request:Request,username:str=Form(...),password:str=Form(...),csrf_token:str=Form('')):
    enforce_csrf(request,csrf_token)
    with conn() as db:
        row=db.execute('SELECT * FROM staff_users WHERE username=? COLLATE NOCASE',(username.strip(),)).fetchone()
    if row and row['active'] and verify_password(password,row['password_hash']):
        request.session.clear()
        request.session.update({'signed_in':True,'role':row['role'],'username':row['username']})
        with conn() as db:db.execute('UPDATE staff_users SET last_login=? WHERE id=?',(now(),row['id']))
        return redirect('/')
    return render(request,'login.html',error='Incorrect username or password')

@app.post('/logout')
def logout(request:Request,csrf_token:str=Form('')):
    enforce_csrf(request,csrf_token)
    request.session.clear(); return redirect('/login')

@app.get('/',response_class=HTMLResponse)
def home(request:Request):
    if not authorized(request): return redirect('/login')
    with conn() as db:
        cases=db.execute('SELECT * FROM cases ORDER BY id DESC LIMIT 30').fetchall()
        stats={k:db.execute(q).fetchone()[0] for k,q in {'cases':'SELECT count(*) FROM cases','open':'SELECT count(*) FROM cases WHERE status NOT IN ("Completed","Cancelled")','policies':'SELECT count(*) FROM policies','pending':'SELECT count(*) FROM cases WHERE status IN ("New","Pending Dispatch")'}.items()}
    with conn() as db:
        stats['vehicles']=db.execute('SELECT count(*) FROM vehicle_production').fetchone()[0] if has_table(db,'vehicle_production') else 0
        stats['completed_tows']=db.execute("SELECT count(*) FROM tow_events WHERE service_status='Completed'").fetchone()[0] if has_table(db,'tow_events') else 0
        stats['imports']=db.execute('SELECT count(*) FROM import_audit').fetchone()[0] if has_table(db,'import_audit') else 0
        trend=[{'day':r['day'],'count':r['cnt']} for r in db.execute("SELECT substr(created_at,1,10) day, count(*) cnt FROM cases GROUP BY substr(created_at,1,10) ORDER BY day DESC LIMIT 14")]
        stats['staff']=db.execute('SELECT count(*) FROM staff_users WHERE active=1').fetchone()[0]
    return render(request,'index.html',cases=cases,stats=stats,trend=list(reversed(trend)))

@app.post('/cases')
def create_case(request:Request,client_name:str=Form(...),phone:str=Form(''),plate:str=Form(''),location:str=Form(''),case_type:str=Form('Towing'),notes:str=Form(''),csrf_token:str=Form('')):
    if not authorized(request): raise HTTPException(401)
    enforce_csrf(request,csrf_token)
    if not client_name.strip(): raise HTTPException(400,'Client name required')
    with conn() as db:
        cur=db.execute('INSERT INTO cases(created_at,client_name,phone,plate,location,case_type,status,notes,provider,eta) VALUES(?,?,?,?,?,?,?,?,?,?)',(now(),client_name.strip(),phone,plate,location,case_type,'New',notes,'',''))
        db.execute('INSERT INTO case_events(case_id,created_at,actor,event) VALUES(?,?,?,?)',(cur.lastrowid,now(),current_actor(request),'Case created'))
    return redirect('/')

STATUS_CHOICES=['New','Pending Dispatch','Assigned','En Route','Arrived','Completed','Escalated','Cancelled']
@app.get('/cases/{case_id}',response_class=HTMLResponse)
def case_detail(request:Request,case_id:int):
    if not authorized(request): return redirect('/login')
    with conn() as db:
        case=db.execute('SELECT * FROM cases WHERE id=?',(case_id,)).fetchone()
        events=db.execute('SELECT * FROM case_events WHERE case_id=? ORDER BY id DESC',(case_id,)).fetchall()
    if not case: raise HTTPException(404)
    return render(request,'case.html',case=case,events=events,status_choices=STATUS_CHOICES)
@app.post('/cases/{case_id}/update')
def case_update(request:Request,case_id:int,status:str=Form(...),provider:str=Form(''),eta:str=Form(''),note:str=Form(''),csrf_token:str=Form('')):
    if not authorized(request): raise HTTPException(401)
    enforce_csrf(request,csrf_token)
    if status not in STATUS_CHOICES: raise HTTPException(400,'Invalid status')
    with conn() as db:
        if not db.execute('SELECT id FROM cases WHERE id=?',(case_id,)).fetchone(): raise HTTPException(404)
        db.execute('UPDATE cases SET status=?,provider=?,eta=? WHERE id=?',(status,provider,eta,case_id))
        db.execute('INSERT INTO case_events(case_id,created_at,actor,event) VALUES(?,?,?,?)',(case_id,now(),current_actor(request),f'Status: {status}; Provider: {provider}; ETA: {eta}; Note: {note}'))
    return redirect(f'/cases/{case_id}')

@app.get('/policies',response_class=HTMLResponse)
def policies(request:Request,q:str=''):
    if not authorized(request): return redirect('/login')
    with conn() as db:
        rows=db.execute('SELECT * FROM policies WHERE policy_number LIKE ? OR insured_name LIKE ? OR plate LIKE ? ORDER BY id DESC LIMIT 100',tuple('%'+q+'%' for _ in range(3))).fetchall()
    return render(request,'policies.html',policies=rows,q=q)
@app.post('/policies')
def add_policy(request:Request,policy_number:str=Form(...),insured_name:str=Form(''),phone:str=Form(''),plate:str=Form(''),start_date:str=Form(''),end_date:str=Form(''),towing_entitlement:str=Form('Review'),status:str=Form('Unknown'),csrf_token:str=Form('')):
    if not authorized(request): raise HTTPException(401)
    enforce_csrf(request,csrf_token)
    with conn() as db:
        db.execute('INSERT INTO policies(policy_number,insured_name,phone,plate,start_date,end_date,towing_entitlement,status) VALUES(?,?,?,?,?,?,?,?) ON CONFLICT(policy_number) DO UPDATE SET insured_name=excluded.insured_name,phone=excluded.phone,plate=excluded.plate,start_date=excluded.start_date,end_date=excluded.end_date,towing_entitlement=excluded.towing_entitlement,status=excluded.status',(policy_number,insured_name,phone,plate,start_date,end_date,towing_entitlement,status))
    return redirect('/policies')

HEADERS=['policy_number','insured_name','phone','plate','start_date','end_date','towing_entitlement','status']
@app.post('/policies/import')
async def import_policies(request:Request,file:UploadFile=File(...),csrf_token:str=Form('')):
    if not authorized(request): raise HTTPException(401)
    enforce_csrf(request,csrf_token)
    data=await file.read()
    if len(data)>5_000_000: raise HTTPException(413,'File too large (5 MB maximum)')
    fname=(file.filename or '').lower()
    if fname.endswith('.xlsx'):
        try:
            book=load_workbook(io.BytesIO(data),read_only=True,data_only=True)
            sheet=book.active
            iterator=sheet.iter_rows(values_only=True)
        except Exception: raise HTTPException(400,'Unable to read XLSX file')
    elif fname.endswith('.csv'):
        try: iterator=csv.reader(io.StringIO(data.decode('utf-8-sig')))
        except Exception: raise HTTPException(400,'Unable to read CSV file')
    else: raise HTTPException(400,'Only XLSX and CSV are supported')
    headers=[str(x).strip().lower() if x is not None else '' for x in next(iterator,[])]
    if 'policy_number' not in headers: raise HTTPException(400,'Missing policy_number header. Download the sample template and try again.')
    positions={h:headers.index(h) for h in HEADERS if h in headers}
    parsed=[]
    for row in iterator:
        if len(parsed)>=10000: raise HTTPException(400,'Import limited to 10,000 records')
        d={h:str(row[i]).strip() if i<len(row) and row[i] is not None else '' for h,i in positions.items()}
        if d.get('policy_number'): parsed.append(d)
    with conn() as db:
        for d in parsed:
            vals=[d.get(h,'') for h in HEADERS]
            db.execute('INSERT INTO policies('+','.join(HEADERS)+') VALUES('+','.join('?' for _ in HEADERS)+') ON CONFLICT(policy_number) DO UPDATE SET '+','.join(f'{h}=excluded.{h}' for h in HEADERS[1:]),vals)
    return redirect('/policies')
@app.get('/health')
def health(): return {'status':'ok','app':'MEARCO Assist demo'}


# Customer requests are restricted by default. Before enabling, configure a durable
# data volume, staffed monitoring, incident response and privacy controls.
REQUEST_LIMIT={}   # process-local abuse friction; use edge/WAF rate limits in production

def limit_request(request):
    ip=request.client.host if request.client else 'unknown'
    now_t=time.monotonic()
    hits=[t for t in REQUEST_LIMIT.get(ip,[]) if now_t-t<600]
    if len(hits)>=5: raise HTTPException(429,'Too many requests. Please contact MEARCO directly.')
    hits.append(now_t); REQUEST_LIMIT[ip]=hits

def validate_fields(name,phone,plate,location,notes):
    if not 2<=len(name.strip())<=120: raise HTTPException(400,'Please enter your name (2–120 characters)')
    if not 6<=len(phone.strip())<=30: raise HTTPException(400,'Please enter a valid contact number')
    if len(plate)>40 or len(location)>350 or len(notes)>2500: raise HTTPException(400,'Request field too long')

# Safe, stand-alone demo: accepts only server-generated fictional records.
# It is available even while public customer intake stays disabled.
@app.get('/demo/assist', response_class=HTMLResponse)
def demo_assist(request:Request):
    return render(request,'demo_assist.html')

@app.post('/demo/assist')
def demo_assist_submit(request:Request,case_type:str=Form('Towing'),area:str=Form('Jbeil'),csrf_token:str=Form('')):
    enforce_csrf(request,csrf_token)
    if case_type not in ('Towing','Motor accident','Complaint','Claim follow-up'):
        raise HTTPException(400,'Invalid demo request type')
    if area not in ('Jbeil','Jounieh','Beirut','Batroun'):
        raise HTTPException(400,'Invalid demo location')
    limit_request(request)
    with conn() as db:
        cur=db.execute('INSERT INTO cases(created_at,client_name,phone,plate,location,case_type,status,notes,provider,eta) VALUES(?,?,?,?,?,?,?,?,?,?)',
            (now(),'DEMO CLIENT — TEST ONLY','00000000','DEMO-0001',area,case_type,'New','DEMO RECORD: no real customer, do not dispatch','',''))
        cid=cur.lastrowid
        db.execute('INSERT INTO case_events(case_id,created_at,actor,event) VALUES(?,?,?,?)',
            (cid,now(),'Demo customer','TEST ONLY — dummy assistance request created'))
    return render(request,'demo_submitted.html',case_id=cid)

@app.get('/assist',response_class=HTMLResponse)
def assist(request:Request):
    return render(request,'assist.html',intake_enabled=PUBLIC_INTAKE_ENABLED)

@app.post('/assist',response_class=HTMLResponse)
def customer_request(request:Request,client_name:str=Form(...),phone:str=Form(...),plate:str=Form(''),location:str=Form(''),case_type:str=Form('Towing'),notes:str=Form(''),is_safe:str=Form(''),website:str=Form(''),csrf_token:str=Form('')):
    if not PUBLIC_INTAKE_ENABLED: raise HTTPException(503,'Online intake is not activated. Call MEARCO for help.')
    enforce_csrf(request,csrf_token)
    if website: raise HTTPException(400,'Invalid request')
    limit_request(request)
    validate_fields(client_name,phone,plate,location,notes)
    if case_type not in ['Towing','Motor accident','Complaint','Claim follow-up']: raise HTTPException(400,'Unsupported request')
    if is_safe not in ['Yes','No','Unsure']: raise HTTPException(400,'Please answer the safety question')
    with conn() as db:
        cur=db.execute('INSERT INTO cases(created_at,client_name,phone,plate,location,case_type,status,notes,provider,eta) VALUES(?,?,?,?,?,?,?,?,?,?)',(now(),client_name.strip(),phone.strip(),plate.strip(),location.strip(),case_type,'Escalated' if is_safe!='Yes' else 'New',f'SAFETY: {is_safe}. {notes.strip()}','',''))
        case_id=cur.lastrowid
        token=secrets.token_urlsafe(24)
        db.execute('INSERT INTO public_tokens(case_id,token_hash) VALUES(?,?)',(case_id,hashlib.sha256(token.encode()).hexdigest()))
        db.execute('INSERT INTO case_events(case_id,created_at,actor,event) VALUES(?,?,?,?)',(case_id,now(),'Client','Online assistance request submitted; safety: '+is_safe))
    return render(request,'submitted.html',case_id=case_id,tracking_token=token,tracking_enabled=PUBLIC_TRACKING_ENABLED)

@app.get('/track',response_class=HTMLResponse)
def track_page(request:Request,case:str='',token:str=''):
    status=None; invalid=False
    if case or token:
        try: cid=int(case.replace('MA-',''))
        except (TypeError,ValueError): cid=0
        if not PUBLIC_TRACKING_ENABLED or not cid or len(token)<20: invalid=True
        else:
            with conn() as db:
                row=db.execute('SELECT token_hash FROM public_tokens WHERE case_id=?',(cid,)).fetchone()
                if not row or not hmac.compare_digest(row['token_hash'],hashlib.sha256(token.encode()).hexdigest()): invalid=True
                else:
                    result=db.execute('SELECT id,status,case_type,eta FROM cases WHERE id=?',(cid,)).fetchone()
                    status=dict(result) if result else None
    return render(request,'track.html',status=status,invalid=invalid,tracking_enabled=PUBLIC_TRACKING_ENABLED)

@app.get('/provider/login',response_class=HTMLResponse)
def provider_login_page(request:Request): return render(request,'provider_login.html',error=None)

@app.post('/provider/login')
def provider_login(request:Request,password:str=Form(...),csrf_token:str=Form('')):
    enforce_csrf(request,csrf_token)
    if PROVIDER_PASSWORD and secrets.compare_digest(password,PROVIDER_PASSWORD):
        request.session['signed_in']=True;request.session['role']='provider'
        return redirect('/provider')
    return templates.TemplateResponse(request,'provider_login.html',{'csrf_token':new_csrf(request),'logged_in':False,'provider_logged_in':False,'error':'Incorrect credentials'},status_code=401)

@app.get('/provider',response_class=HTMLResponse)
def provider_jobs(request:Request):
    if not provider_authorized(request): return redirect('/provider/login')
    with conn() as db:
        rows=db.execute("SELECT c.*,p.assigned_driver,p.provider_note FROM cases c LEFT JOIN provider_jobs p ON p.case_id=c.id WHERE c.provider != '' AND c.case_type='Towing' AND c.status NOT IN ('Completed','Cancelled') ORDER BY c.id DESC LIMIT 100").fetchall()
    return render(request,'provider.html',jobs=rows)

@app.post('/provider/jobs/{case_id}/update')
def provider_job_update(request:Request,case_id:int,status:str=Form(...),driver:str=Form(''),note:str=Form(''),csrf_token:str=Form('')):
    if not provider_authorized(request): raise HTTPException(401)
    enforce_csrf(request,csrf_token)
    if status not in ['Assigned','En Route','Arrived','Completed']: raise HTTPException(400,'Status not allowed')
    with conn() as db:
        case=db.execute("SELECT id FROM cases WHERE id=? AND provider!='' AND case_type='Towing'",(case_id,)).fetchone()
        if not case: raise HTTPException(404)
        db.execute('UPDATE cases SET status=? WHERE id=?',(status,case_id))
        db.execute('INSERT INTO provider_jobs(case_id,assigned_driver,provider_note,updated_at) VALUES(?,?,?,?) ON CONFLICT(case_id) DO UPDATE SET assigned_driver=excluded.assigned_driver,provider_note=excluded.provider_note,updated_at=excluded.updated_at',(case_id,driver[:120],note[:1000],now()))
        db.execute('INSERT INTO case_events(case_id,created_at,actor,event) VALUES(?,?,?,?)',(case_id,now(),'Towing provider',f'{status}; driver: {driver[:120]}; note: {note[:1000]}'))
    return redirect('/provider')


# Browser-based conversation demo. Never accepts genuine personal information.
# Only synthetic records are created in the employee dashboard.
@app.get('/chat',response_class=HTMLResponse)
def chat_demo(request:Request):
    return render(request,'chat.html')

@app.post('/api/chat-demo/submit')
async def chat_demo_submit(request:Request):
    try:
        data=await request.json()
    except Exception:
        raise HTTPException(400,'Invalid request')
    enforce_csrf(request,str(data.get('csrf_token','')))
    case_type=data.get('case_type')
    area=data.get('area')
    issue=data.get('issue')
    safe=data.get('safe')
    if case_type not in ('Towing','Motor accident','Complaint','Claim follow-up'):
        raise HTTPException(400,'Invalid category')
    if area not in ('Jbeil','Jounieh','Beirut','Batroun'):
        raise HTTPException(400,'Invalid area')
    if issue not in ('Breakdown','Flat tire','Collision','Other'):
        raise HTTPException(400,'Invalid scenario')
    if safe not in ('Yes','No','Unsure'):
        raise HTTPException(400,'Invalid safety answer')
    limit_request(request)
    with conn() as db:
        cur=db.execute('INSERT INTO cases(created_at,client_name,phone,plate,location,case_type,status,notes,provider,eta) VALUES(?,?,?,?,?,?,?,?,?,?)',
            (now(),'DEMO CHAT CLIENT — NOT REAL','00000000','DEMO-CHAT',area,case_type,'Escalated' if safe!='Yes' else 'New',
             f'TEST CHAT — NO REAL REQUEST. Scenario: {issue}. Safe: {safe}. Do not dispatch.','',''))
        cid=cur.lastrowid
        db.execute('INSERT INTO case_events(case_id,created_at,actor,event) VALUES(?,?,?,?)',
           (cid,now(),'Website chat demo','TEST ONLY — fictional chatbot request. Do not dispatch.'))
    return {'case_ref':f'MA-{cid:06d}','demo':True}


# === v0.5 vehicle production and towing ledger (staff-only) ===
from collections import Counter
import unicodedata
from datetime import date as date_type
from openpyxl.utils.datetime import from_excel

def norm_identifier(v):
    s=unicodedata.normalize('NFKC',str(v or '')).upper().strip()
    return ''.join(c for c in s if c.isalnum())
def valstr(v):
    if v is None:return ''
    if isinstance(v,(datetime,date_type)):return v.strftime('%Y-%m-%d')
    return str(v).strip()
def iso_date(v):
    if not v:return ''
    if isinstance(v,(datetime,date_type)):return v.strftime('%Y-%m-%d')
    if isinstance(v,(int,float)):
        try:return from_excel(v).strftime('%Y-%m-%d')
        except Exception:return ''
    s=valstr(v).split(' ')[0]
    for fmt in ('%Y-%m-%d','%d/%m/%Y','%m/%d/%Y','%d-%m-%Y','%d/%m/%y','%d-%m-%y'):
        try:return datetime.strptime(s,fmt).strftime('%Y-%m-%d')
        except ValueError: pass
    return ''

ALIASES={
'policy_number':['policy number','policy no','policy no.','policy_number','policy','pol no','pol number','policy nbr','contract no','policyno'],
'insured_name':['insured name','insured','client name','client','name','assured name','policyholder','subscriber'],
'phone':['mobile','phone','telephone','tel','mobile number','phone number','contact'],
'plate':['plate','plate number','plate no','car plate','vehicle plate','registration','registration number','platenumber','matricule','plate no.'],
'chassis':['chassis','chassis number','chassis no','chassis no.','vin','vin number','chassis_number','vehicle identification number'],
'start_date':['inception','inception date','effective date','from','start date','start_date','date from','policy from','date of inception'],
'end_date':['expiry','expiry date','expiration date','to','end date','end_date','date to','policy to','date of expiry'],
'limit':['towing limit','towing_limit','max tows','allowed tows','allowed_tows','towing times','towing entitlement count','tows allowed','number of tows','towing quota','towing limit per year'],
'status':['status','policy status'],
'event_id':['tow id','towing id','reference','tow reference','job id','service id','dispatch id','case number','request id'],
'event_date':['towing date','tow date','date','service date','date of towing','date of service','assistance date'],
'from_location':['from location','from_location','pickup','pickup location','from','departure','origin','location from'],
'to_location':['to location','to_location','destination','to','arrival','location to'],
'provider':['provider','towing provider','tow company','supplier','contractor'],
'service_status':['service status','tow status','job status','towing status','status','service_status'],
'count':['tow count','towing count','number of tows','number of times','times','quantity','qty','count'],
}
def clean_header(v):
    s=unicodedata.normalize('NFKC',valstr(v)).lower()
    return re.sub(r'[^a-z0-9]+',' ',s).strip()
def mapping(header,fields):
    h=[clean_header(x) for x in header]; out={}
    for f in fields:
        for alias in ALIASES[f]:
            if clean_header(alias) in h:
                out[f]=h.index(clean_header(alias));break
    return out

def ensure_ledger_tables():
    with conn() as db:
        db.execute('CREATE TABLE IF NOT EXISTS vehicle_production(id INTEGER PRIMARY KEY, policy_number TEXT NOT NULL, insured_name TEXT, phone TEXT, plate TEXT, chassis TEXT, plate_key TEXT, chassis_key TEXT, start_date TEXT, end_date TEXT, tow_limit INTEGER, status TEXT, imported_at TEXT, UNIQUE(policy_number,plate_key,chassis_key,start_date))')
        db.execute('CREATE TABLE IF NOT EXISTS tow_events(id INTEGER PRIMARY KEY, event_key TEXT UNIQUE NOT NULL, event_id TEXT, event_date TEXT, plate TEXT, chassis TEXT, plate_key TEXT, chassis_key TEXT, from_location TEXT, to_location TEXT, provider TEXT, source_name TEXT, imported_at TEXT)')
        db.execute('CREATE TABLE IF NOT EXISTS import_audit(id INTEGER PRIMARY KEY, imported_at TEXT, category TEXT, filename TEXT, rows_read INTEGER, accepted INTEGER, skipped INTEGER, message TEXT)')
        cols={row[1] for row in db.execute('PRAGMA table_info(tow_events)')}
        if 'service_status' not in cols:db.execute("ALTER TABLE tow_events ADD COLUMN service_status TEXT DEFAULT 'Completed'")
        if 'recorded_by' not in cols:db.execute("ALTER TABLE tow_events ADD COLUMN recorded_by TEXT DEFAULT 'Historical import'")
ensure_ledger_tables()

def read_tabular(data, filename):
    if len(data)>12_000_000:raise HTTPException(413,'Maximum file size: 12 MB')
    if filename.lower().endswith('.xlsx'):
        try:
            wb=load_workbook(io.BytesIO(data),data_only=True,read_only=True)
            iterator=wb.active.iter_rows(values_only=True)
            result=[]
            for i,row in enumerate(iterator):
                if i>30001:raise HTTPException(400,'Maximum 30,000 data rows per import')
                result.append(row)
            wb.close();return result
        except HTTPException:raise
        except Exception as exc:raise HTTPException(400,'Unable to read XLSX spreadsheet') from exc
    if filename.lower().endswith('.csv'):
        try:
            reader=csv.reader(io.StringIO(data.decode('utf-8-sig')))
            result=[]
            for i,row in enumerate(reader):
                if i>30001:raise HTTPException(400,'Maximum 30,000 data rows per import')
                result.append(row)
            return result
        except HTTPException:raise
        except Exception as exc:raise HTTPException(400,'Unable to read UTF-8 CSV') from exc
    raise HTTPException(400,'Upload XLSX or CSV only')

def event_identity(ref, event_date,plate_key,chassis_key,from_loc,to_loc):
    # Stable cross-upload key. Job ID preferred; otherwise exact detail-based fingerprint.
    base='REF:'+norm_identifier(ref) if ref else 'DETAIL:'+'|'.join((event_date,plate_key,chassis_key,norm_identifier(from_loc),norm_identifier(to_loc)))
    return hashlib.sha256(base.encode('utf-8')).hexdigest()

@app.get('/vehicle-usage',response_class=HTMLResponse)
def vehicle_usage(request:Request,q:str='',message:str=''):
    if not authorized(request):return redirect('/login')
    q=q.strip()[:120]
    with conn() as db:
        prods=[dict(r) for r in db.execute('SELECT * FROM vehicle_production ORDER BY id DESC LIMIT 30000')]
        tows=[dict(r) for r in db.execute('SELECT * FROM tow_events ORDER BY id DESC LIMIT 30000')]
        audits=db.execute('SELECT * FROM import_audit ORDER BY id DESC LIMIT 12').fetchall()
    by_plate={};by_chassis={}
    for p in prods:
        if p['plate_key']:by_plate.setdefault(p['plate_key'],set()).add(p['id'])
        if p['chassis_key']:by_chassis.setdefault(p['chassis_key'],set()).add(p['id'])
    counts=Counter(); ambiguous=[]; unmatched=[]; undated=[]; outside=[]
    for t in tows:
        if t.get('service_status') not in ('Completed', None, ''):
            continue
        chassis_hits=by_chassis.get(t['chassis_key'],set()) if t['chassis_key'] else set()
        plate_hits=by_plate.get(t['plate_key'],set()) if t['plate_key'] else set()
        candidates=chassis_hits if chassis_hits else plate_hits
        if chassis_hits and plate_hits and not (chassis_hits & plate_hits):ambiguous.append(t);continue
        if not candidates:unmatched.append(t);continue
        # Multiple policies for a vehicle: allocate by event date and policy period.
        viable=[]
        for p in prods:
            if p['id'] not in candidates:continue
            d=t['event_date']
            if d and (not p['start_date'] or p['start_date']<=d) and (not p['end_date'] or d<=p['end_date']):viable.append(p)
        if not t['event_date']:undated.append(t);continue
        if len(viable)==1:counts[viable[0]['id']]+=1
        elif len(viable)>1:ambiguous.append(t)
        else:outside.append(t)
    rows=[]
    for p in prods:
        used=counts[p['id']]
        limit=p['tow_limit']
        remaining=max(limit-used,0) if limit is not None else None
        row={**p,'used':used,'remaining':remaining,'exceeded':limit is not None and used>limit}
        if not q or q.lower() in (' '.join(valstr(row.get(k,'')) for k in ('policy_number','insured_name','plate','chassis','phone'))).lower():rows.append(row)
    stats={'policies':len(prods),'tows':sum(t.get('service_status') in ('Completed',None,'') for t in tows),'pending_tows':sum(t.get('service_status') not in ('Completed',None,'','Cancelled') for t in tows),'cancelled_tows':sum(t.get('service_status')=='Cancelled' for t in tows),'matched':sum(counts.values()),'unmatched':len(unmatched),'ambiguous':len(ambiguous),'undated':len(undated),'outside':len(outside),'unknown_limit':sum(p['tow_limit'] is None for p in prods)}
    return render(request,'vehicle_usage.html',rows=rows[:500],stats=stats,q=q,message=message,audits=audits,unmatched=unmatched[:30],ambiguous=ambiguous[:30],undated=undated[:30],outside=outside[:30])

@app.post('/vehicle-usage/import')
async def vehicle_import(request:Request,category:str=Form(...),file:UploadFile=File(...),csrf_token:str=Form('')):
    if not authorized(request):raise HTTPException(401)
    enforce_csrf(request,csrf_token)
    if category not in ('production','tows'):raise HTTPException(400,'Unknown import category')
    data=await file.read();filename=os.path.basename(file.filename or 'upload')[:150]
    rows=read_tabular(data,filename)
    if not rows:raise HTTPException(400,'Spreadsheet is empty')
    fields=['policy_number','insured_name','phone','plate','chassis','start_date','end_date','limit','status'] if category=='production' else ['event_id','event_date','plate','chassis','from_location','to_location','provider','service_status']
    indexes=mapping(rows[0],fields)
    if category=='production' and not indexes.get('policy_number') is not None:raise HTTPException(400,'No policy number column found. Use the downloadable sample template.')
    if category=='production' and 'plate' not in indexes and 'chassis' not in indexes:raise HTTPException(400,'Must include plate or chassis column.')
    if category=='tows' and ('plate' not in indexes and 'chassis' not in indexes):raise HTTPException(400,'Towing log needs plate or chassis for matching.')
    if category=='tows' and 'event_date' not in indexes:raise HTTPException(400,'Towing log needs a tow/service date to allocate usage to policy terms.')
    def get(row,key):
        i=indexes.get(key)
        return row[i] if i is not None and i<len(row) else None
    parsed=[]; skipped=0
    for row in rows[1:]:
        if not any(x is not None and str(x).strip() for x in row):continue
        plate=valstr(get(row,'plate'));chassis=valstr(get(row,'chassis'))
        pk=norm_identifier(plate);ck=norm_identifier(chassis)
        if category=='production':
            pol=valstr(get(row,'policy_number')); start=iso_date(get(row,'start_date'));end=iso_date(get(row,'end_date'))
            raw_limit=get(row,'limit');limit=None
            if raw_limit is not None and valstr(raw_limit):
                try:
                    v=float(raw_limit)
                    if not 0<=v<=100 or not v.is_integer():raise ValueError()
                    limit=int(v)
                except (ValueError,TypeError):skipped+=1;continue
            if not pol or (not pk and not ck):skipped+=1;continue
            parsed.append((pol,valstr(get(row,'insured_name')),valstr(get(row,'phone')),plate,chassis,pk,ck,start,end,limit,valstr(get(row,'status')),now()))
        else:
            d=iso_date(get(row,'event_date'))
            if not (pk or ck) or not d:skipped+=1;continue
            ref=valstr(get(row,'event_id'));fr=valstr(get(row,'from_location'));to=valstr(get(row,'to_location'))
            key=event_identity(ref,d,pk,ck,fr,to)
            status=valstr(get(row,'service_status')) or 'Completed'
            if status.lower() not in ('completed','cancelled','pending','assigned','en route','arrived'):skipped+=1;continue
            status=next(x for x in ('Completed','Cancelled','Pending','Assigned','En Route','Arrived') if x.lower()==status.lower())
            parsed.append((key,ref,d,plate,chassis,pk,ck,fr,to,valstr(get(row,'provider')),filename,now(),status,'Excel import'))
    with conn() as db:
        if category=='production':
            # Repeated imports update the same policy+vehicle+term without multiplying vehicles.
            for p in parsed:
                db.execute('INSERT INTO vehicle_production(policy_number,insured_name,phone,plate,chassis,plate_key,chassis_key,start_date,end_date,tow_limit,status,imported_at) VALUES(?,?,?,?,?,?,?,?,?,?,?,?) ON CONFLICT(policy_number,plate_key,chassis_key,start_date) DO UPDATE SET insured_name=excluded.insured_name,phone=excluded.phone,plate=excluded.plate,chassis=excluded.chassis,end_date=excluded.end_date,tow_limit=excluded.tow_limit,status=excluded.status,imported_at=excluded.imported_at',p)
        else:
            for t in parsed:
                db.execute('INSERT INTO tow_events(event_key,event_id,event_date,plate,chassis,plate_key,chassis_key,from_location,to_location,provider,source_name,imported_at,service_status,recorded_by) VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?,?) ON CONFLICT(event_key) DO UPDATE SET provider=excluded.provider,source_name=excluded.source_name,imported_at=excluded.imported_at,service_status=excluded.service_status',t)
        db.execute('INSERT INTO import_audit(imported_at,category,filename,rows_read,accepted,skipped,message) VALUES(?,?,?,?,?,?,?)',(now(),category,filename,len(rows)-1,len(parsed),skipped,'Rows processed; duplicates updated rather than duplicated'))
    return redirect('/vehicle-usage?message='+('Production' if category=='production' else 'Towing')+'+import+processed')


# === v0.6 quick towing entry and internal alerts ===
@app.post('/vehicle-usage/record-tow')
def record_tow(request:Request, event_id:str=Form(...),event_date:str=Form(...),plate:str=Form(''),chassis:str=Form(''),from_location:str=Form(''),to_location:str=Form(''),provider:str=Form(''),service_status:str=Form('Completed'),csrf_token:str=Form('')):
    if not authorized(request):raise HTTPException(401)
    enforce_csrf(request,csrf_token)
    if service_status not in ('Completed','Cancelled','Pending','Assigned','En Route','Arrived'):raise HTTPException(400,'Invalid tow status')
    event_id=event_id.strip()[:100]
    if not event_id:raise HTTPException(400,'Unique dispatch reference required')
    d=iso_date(event_date)
    pk,ck=norm_identifier(plate),norm_identifier(chassis)
    if not d or not (pk or ck):raise HTTPException(400,'Date and plate or chassis required')
    key=event_identity(event_id,d,pk,ck,from_location,to_location)
    with conn() as db:
        exists=db.execute('SELECT id FROM tow_events WHERE event_key=?',(key,)).fetchone()
        if exists:
            db.execute('UPDATE tow_events SET service_status=?,provider=?,from_location=?,to_location=?,imported_at=?,recorded_by=? WHERE id=?',(service_status,provider[:120],from_location[:250],to_location[:250],now(),current_actor(request),exists['id']))
        else:
            db.execute('INSERT INTO tow_events(event_key,event_id,event_date,plate,chassis,plate_key,chassis_key,from_location,to_location,provider,source_name,imported_at,service_status,recorded_by) VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?,?)',(key,event_id,d,plate[:120],chassis[:120],pk,ck,from_location[:250],to_location[:250],provider[:120],'Manual entry',now(),service_status,current_actor(request)))
    return redirect('/vehicle-usage?message=Tow+record+saved')

@app.get('/api/operator/pending')
def operator_pending(request:Request):
    if not authorized(request):raise HTTPException(401)
    with conn() as db:
        rows=db.execute("SELECT id,created_at,case_type,location FROM cases WHERE status IN ('New','Pending Dispatch','Escalated') ORDER BY id DESC LIMIT 30").fetchall()
    return {'pending':[dict(r) for r in rows],'count':len(rows)}

# === v0.7 staff administration: owner/admin controls users; operators never manage accounts ===
@app.get('/users',response_class=HTMLResponse)
def user_list(request:Request,message:str=''):
    if not is_admin(request):raise HTTPException(403,'Administrator only')
    with conn() as db:
        users=db.execute('SELECT id,username,display_name,role,active,created_at,last_login FROM staff_users ORDER BY username').fetchall()
        audit=db.execute('SELECT * FROM staff_audit ORDER BY id DESC LIMIT 15').fetchall()
    return render(request,'users.html',users=users,message=message,audit=audit)

@app.post('/users/create')
def user_create(request:Request,username:str=Form(...),display_name:str=Form(...),role:str=Form(...),password:str=Form(...),csrf_token:str=Form('')):
    if not is_admin(request):raise HTTPException(403)
    enforce_csrf(request,csrf_token)
    username=username.strip().lower()
    if not re.fullmatch(r'[a-z0-9._-]{3,40}',username):raise HTTPException(400,'Username: 3–40 letters, numbers, dots, hyphens or underscores')
    if role not in ('operator','supervisor'):raise HTTPException(400,'Invalid role')
    if not (2<=len(display_name.strip())<=80) or len(password)<12 or len(password)>200:raise HTTPException(400,'Display name required and password must be 12–200 characters')
    try:
        with conn() as db:
            db.execute('INSERT INTO staff_users(username,display_name,role,active,password_hash,created_at) VALUES(?,?,?,?,?,?)',(username,display_name.strip(),role,1,hash_password(password),now()))
            db.execute('INSERT INTO staff_audit(created_at,actor,action,subject) VALUES(?,?,?,?)',(now(),current_actor(request),'Create staff user',username))
    except sqlite3.IntegrityError:raise HTTPException(409,'Username already exists')
    return redirect('/users?message=User+created')

@app.post('/users/{user_id}/status')
def staff_status(request:Request,user_id:int,active:str=Form(...),csrf_token:str=Form('')):
    if not is_admin(request):raise HTTPException(403)
    enforce_csrf(request,csrf_token)
    if active not in ('0','1'):raise HTTPException(400)
    with conn() as db:
        user=db.execute('SELECT username FROM staff_users WHERE id=?',(user_id,)).fetchone()
        if not user:raise HTTPException(404)
        db.execute('UPDATE staff_users SET active=? WHERE id=?',(int(active),user_id))
        db.execute('INSERT INTO staff_audit(created_at,actor,action,subject) VALUES(?,?,?,?)',(now(),current_actor(request),'Enable staff' if active=='1' else 'Disable staff',user['username']))
    return redirect('/users?message=Account+updated')

@app.post('/users/{user_id}/reset')
def staff_reset(request:Request,user_id:int,password:str=Form(...),csrf_token:str=Form('')):
    if not is_admin(request):raise HTTPException(403)
    enforce_csrf(request,csrf_token)
    if not 12<=len(password)<=200:raise HTTPException(400,'Password must be at least 12 characters')
    with conn() as db:
        user=db.execute('SELECT username FROM staff_users WHERE id=?',(user_id,)).fetchone()
        if not user:raise HTTPException(404)
        db.execute('UPDATE staff_users SET password_hash=? WHERE id=?',(hash_password(password),user_id))
        db.execute('INSERT INTO staff_audit(created_at,actor,action,subject) VALUES(?,?,?,?)',(now(),current_actor(request),'Reset password',user['username']))
    return redirect('/users?message=Password+updated')
